"""Streamlit UI for the JPC-gated dynamic checklist upload pipeline."""
from __future__ import annotations

import io
import os

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app_pipeline import process_upload_batch
from central_config import build_central_adapter
from central_write_adapter import WriteStatus
from metadata_engine import load_checklist_departments
from photo_engine import build_report, load_mapping, load_raw_data


BASE = os.path.dirname(__file__)
DASH = os.path.join(BASE, "master", "PDI_Dashboard_QC.xlsx")
try:
    _streamlit_secrets = st.secrets
except Exception:
    _streamlit_secrets = {}
CENTRAL_ADAPTER = build_central_adapter(_streamlit_secrets)


@st.cache_data
def load_report_data():
    mapping, conflicts = load_mapping(DASH) if os.path.exists(DASH) else ({}, {})
    raw = load_raw_data(DASH) if os.path.exists(DASH) else {}
    return mapping, conflicts, raw


def _excel_bytes(report: pd.DataFrame) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Traceability Report"
    for column, heading in enumerate(report.columns, 1):
        sheet.cell(1, column, heading)
    for row_number, row in enumerate(report.itertuples(index=False), 2):
        for column, value in enumerate(row, 1):
            sheet.cell(row_number, column, value)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(
            max(max(len(str(cell.value or "")) for cell in column) + 2, 10), 40)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


st.set_page_config(page_title="PDI Photo Upload – Phase 2", page_icon="📋", layout="centered")
st.title("📋 PDI Photo Upload – Phase 2")
st.caption("JPC validation → Dynamic row/tick detection → Observation text → Batch submission")

mapping, conflicts, raw = load_report_data()
checklist_departments = load_checklist_departments(os.path.join(BASE, "master", "Observation_Checklist.xlsx"))
combined_mapping = dict(checklist_departments)
combined_mapping.update(mapping)

with st.form("upload"):
    inspector = st.text_input("Inspector Name *", placeholder="e.g. Harish")
    entered_jpc = st.text_input("JPC Number *", placeholder="e.g. VH-3009")
    photos = st.file_uploader("Upload checklist photos (all pages) *",
                              type=["jpg", "jpeg", "png"], accept_multiple_files=True)
    submitted = st.form_submit_button("Process Checklist")

if submitted:
    if not inspector.strip():
        st.error("Inspector Name is required.")
        st.stop()
    if not entered_jpc.strip():
        st.error("JPC Number is required.")
        st.stop()
    if not photos:
        st.error("Upload at least one checklist photo.")
        st.stop()

    photo_payloads = [{"filename": photo.name, "data": photo.getvalue()} for photo in photos]
    result = process_upload_batch(inspector, entered_jpc, photo_payloads,
                                  central_adapter=CENTRAL_ADAPTER,
                                  metadata_mapping=combined_mapping)
    st.subheader("Batch JPC validation")
    jpc_rows = [{"Photo": check.filename, "Status": check.status,
                 "Candidates": ", ".join(check.candidates), "Reason": check.reason}
                for check in result.photo_jpc_results]
    if jpc_rows:
        st.dataframe(pd.DataFrame(jpc_rows), use_container_width=True, hide_index=True)
    if result.status == "JPC_REJECTED":
        st.error(result.message)
        st.stop()

    if result.status == "NO_OBSERVATIONS":
        st.info(result.message)
        st.stop()
    if result.status == "REJECTED":
        st.error(result.message)
        st.stop()

    st.subheader("Confirmed NOT OK observations")
    st.dataframe(pd.DataFrame(result.observations), use_container_width=True, hide_index=True)
    report_candidates = pd.DataFrame([
        {"Observation": record["Observation"],
         "OCR Match %": round(float(record["OCR Confidence"]) * 100, 1),
         "OK Ink": 0, "NOT OK Ink": 0, "Detected State": "NOT OK"}
        for record in result.observations
    ])
    report, review = build_report(result.jpc, inspector, report_candidates, raw, mapping)
    st.subheader("Traceability Report Preview")
    st.dataframe(report, use_container_width=True, hide_index=True)
    if not review.empty:
        st.subheader("Review Queue")
        st.dataframe(review, use_container_width=True, hide_index=True)

    if result.status == WriteStatus.SUCCESS.value:
        st.success(f"Central write confirmed: {result.write_result.records_written} record(s).")
        st.download_button("⬇️ Download Traceability Report Excel", _excel_bytes(report),
                           "PDI_Automated_Traceability_Report_Phase2.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    elif result.status == WriteStatus.DUPLICATE.value:
        st.info(f"Central system reports this JPC batch as a duplicate: {result.message}")
    else:
        st.error(f"Central write not confirmed ({result.status}): {result.message}")
        st.caption("The report above is a local preview only; no central-write success is claimed.")
        st.download_button("⬇️ Download local report preview (not submitted)", _excel_bytes(report),
                           "PDI_Automated_Traceability_Report_Phase2.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.caption(f"Inspector: {inspector} | JPC: {result.jpc} | Photos: {len(photos)} | "
               f"Observations: {len(result.observations)}")
    if result.timings_ms:
        st.caption("Processing times (ms): " + ", ".join(
            f"{name}={value:.1f}" for name, value in result.timings_ms.items()))
