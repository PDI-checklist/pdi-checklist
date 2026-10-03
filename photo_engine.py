import io, re, os, tempfile
from datetime import datetime
import cv2, numpy as np, pandas as pd, pytesseract
from rapidfuzz import process, fuzz
from pytesseract import Output
from openpyxl import load_workbook

OBS_HEADERS = ["S.No","JPC Number","Date","kVA","Shift","Phase","Variant","Pole","Observation","Department","Station","Defect Category","Status","Cleared By","Closure Date","Closure Remarks"]


def load_checklist(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    rows=[]
    for ws in wb.worksheets:
        section=None
        for r in range(1, ws.max_row+1):
            a,b=ws.cell(r,1).value,ws.cell(r,2).value
            if isinstance(a,str) and a.strip().upper() in {"INTEGRATION","TESTING","STUFFING","BASE FRAME","POWDER COATING","CUSTOMER COMPLAINT"}:
                section=a.strip()
            if isinstance(a,(int,float)) and b:
                rows.append({"observation":str(b).strip(),"sheet":ws.title,"section":section or ws.title,"sno":int(a)})
    return pd.DataFrame(rows).drop_duplicates(subset=["observation"]).reset_index(drop=True)


def load_mapping(dashboard_path):
    wb=load_workbook(dashboard_path,data_only=True,read_only=True)
    ws=wb["Observation Log"]
    mp={}
    conflicts={}
    for row in ws.iter_rows(min_row=8, max_col=15, values_only=True):
        obs=row[11]
        if not obs: continue
        key=str(obs).strip().lower()
        val=(row[12] or "",row[13] or "",row[14] or "")
        if key in mp and mp[key]!=val:
            conflicts.setdefault(key,set()).update([mp[key],val])
        else:
            mp[key]=val
    return mp, conflicts


def load_raw_data(dashboard_path):
    wb=load_workbook(dashboard_path,data_only=True,read_only=True)
    ws=wb["Raw Data"]
    out={}
    for row in ws.iter_rows(min_row=7,max_col=15,values_only=True):
        jpc=row[6]
        if jpc:
            out[str(jpc).strip().upper()]={"Date":row[4],"Shift":row[5],"kVA":row[7],"Phase":row[9],"Variant":row[10],"Pole":row[11]}
    return out


def ocr_text(img):
    d=pytesseract.image_to_data(img, config="--psm 6", output_type=Output.DICT)
    lines={}
    for i,t in enumerate(d["text"]):
        t=t.strip()
        try: conf=float(d["conf"][i])
        except: conf=0
        if not t or conf<20: continue
        key=(d["block_num"][i],d["par_num"][i],d["line_num"][i])
        lines.setdefault(key,[]).append(i)
    items=[]
    for inds in lines.values():
        text=" ".join(d["text"][i] for i in inds).strip()
        if not text: continue
        items.append({"text":text,"x":min(d["left"][i] for i in inds),"y":int(np.mean([d["top"][i]+d["height"][i]/2 for i in inds])),"w":max(d["left"][i]+d["width"][i] for i in inds)-min(d["left"][i] for i in inds)})
    return items


def header_extract(img):
    text=pytesseract.image_to_string(img, config="--psm 6")
    t=" ".join(text.split())
    jpc=None; kva=None; date=None
    m=re.search(r"JPC\s*(?:NUMBER|NO)?\s*[:\-]?\s*([A-Z0-9]+\s*[-/]\s*[A-Z0-9]+)",t,re.I)
    if m: jpc=m.group(1).replace(" ","").upper()
    m=re.search(r"KVA\s*[:\-]?\s*([0-9]+(?:\.[0-9]+)?)",t,re.I)
    if m: kva=m.group(1)
    m=re.search(r"DATE\s*[:\-]?\s*([0-9]{1,2}\s*[\-/]\s*[0-9]{1,2}\s*[\-/]\s*[0-9]{2,4})",t,re.I)
    if m: date=m.group(1).replace(" ","")
    return {"jpc":jpc,"kva":kva,"date":date,"ocr_text":t}


def vertical_bounds(img):
    gray=cv2.cvtColor(img,cv2.COLOR_BGR2GRAY)
    th=cv2.adaptiveThreshold(gray,255,cv2.ADAPTIVE_THRESH_MEAN_C,cv2.THRESH_BINARY_INV,31,10)
    k=max(45,img.shape[0]//25)
    v=cv2.morphologyEx(th,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_RECT,(1,k)))
    n,lab,stats,cent=cv2.connectedComponentsWithStats(v)
    xs=[]
    for i in range(1,n):
        x,y,w,h,a=stats[i]
        if h>img.shape[0]*0.10 and w<15 and x>img.shape[1]*0.35 and x<img.shape[1]*0.95:
            xs.append(x+w//2)
    xs=sorted(xs); clusters=[]
    for x in xs:
        if not clusters or x-clusters[-1][-1]>15: clusters.append([x])
        else: clusters[-1].append(x)
    return [int(np.mean(c)) for c in clusters]


def ink_score(gray,x1,x2,y1,y2):
    roi=gray[max(0,y1):min(gray.shape[0],y2),max(0,x1):min(gray.shape[1],x2)]
    if roi.size==0: return 0
    bw=cv2.adaptiveThreshold(roi,255,cv2.ADAPTIVE_THRESH_GAUSSIAN_C,cv2.THRESH_BINARY_INV,21,8)
    h=cv2.morphologyEx(bw,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_RECT,(max(8,roi.shape[1]//2),1)))
    v=cv2.morphologyEx(bw,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_RECT,(1,max(8,roi.shape[0]//2))))
    clean=cv2.subtract(bw,cv2.bitwise_or(h,v))
    n,lab,stats,cent=cv2.connectedComponentsWithStats(clean)
    score=0
    for i in range(1,n):
        x,y,w,h,a=stats[i]
        if 3<=a<=roi.shape[0]*roi.shape[1]*0.35 and w>=2 and h>=2:
            score += a
    return int(score)


def detect_candidates(img, checklist_df):
    lines=ocr_text(img)
    choices=checklist_df.observation.tolist()
    gray=cv2.cvtColor(img,cv2.COLOR_BGR2GRAY)
    bounds=vertical_bounds(img)
    xb=[x for x in bounds if 0.45*img.shape[1]<x<0.85*img.shape[1]]
    xb=sorted(xb)
    if len(xb)>=3:
        x1,x2,x3=xb[:3]
    else:
        x1,x2,x3=int(img.shape[1]*.61),int(img.shape[1]*.68),int(img.shape[1]*.75)
    found=[]
        def has_handwritten_tick(cell):
        if cell is None or cell.size == 0:
            return False

        h, w = cell.shape[:2]

        if h > 6 and w > 6:
            cell = cell[3:h-3, 3:w-3]

        edges = cv2.Canny(cell, 40, 120)

        lines_h = cv2.HoughLinesP(
            edges,
            1,
            np.pi / 180,
            threshold=4,
            minLineLength=3,
            maxLineGap=5
        )

        if lines_h is None:
            return False

        positive = 0
        negative = 0

        for ln in lines_h:
            coords = ln[0]
            xa, ya, xb, yb = map(int, coords)

            dx = xb - xa
            dy = yb - ya

            if abs(dx) < 2:
                continue

            length = (dx * dx + dy * dy) ** 0.5
            slope = dy / dx

            if 3 <= length <= 35 and 0.20 <= abs(slope) <= 5:
                if slope > 0:
                    positive += 1
                else:
                    negative += 1

        return positive >= 1 and negative >= 1


    for line in lines:
        if line["x"] > x1:
            continue

        match = process.extractOne(
            line["text"],
            choices,
            scorer=fuzz.token_set_ratio
        )

        if not match or match[1] < 62:
            continue

        obs = match[0]
        conf = float(match[1])

        if obs in used:
            continue

        y = line["y"]

        ok_cell = gray[y-18:y+18, x1+8:x2-8]
        nok_cell = gray[y-18:y+18, x2+8:x3-8]

        ok_tick = has_handwritten_tick(ok_cell)
        nok_tick = has_handwritten_tick(nok_cell)

        if nok_tick and not ok_tick:
            state = "NOT OK"
        elif ok_tick and not nok_tick:
            state = "OK"
        else:
            state = "REVIEW"

        if state in ("NOT OK", "REVIEW"):
            found.append({
                "Observation": obs,
                "OCR Match %": round(conf, 1),
                "OK Ink": int(ok_tick),
                "NOT OK Ink": int(nok_tick),
                "Detected State": state
            })

        used.add(obs)

    return pd.DataFrame(found)


def build_report(jpc, inspector, candidates, raw_lookup, mapping):
    meta=raw_lookup.get(jpc.strip().upper(),{})
    rows=[]; review=[]
    for i,rec in candidates.iterrows():
        obs=rec["Observation"]
        if rec["Detected State"] != "NOT OK":
            if rec["Detected State"] == "REVIEW":
                review.append({
                    "Observation": obs,
                    "Reason": "Low tick confidence",
                    "Detected State": rec["Detected State"],
                    "OCR Match %": rec["OCR Match %"]
                })
            continue
        m=mapping.get(obs.lower())
        if m:
            dept,station,defect=m
            ambiguous=False
        else:
            dept=station=defect=""
            ambiguous=True
         
        if ambiguous:
            review.append({"Observation":obs,"Reason":"No unambiguous mapping in Observation Log","Detected State":rec["Detected State"],"OCR Match %":rec["OCR Match %"]})
            continue
        rows.append([len(rows)+1,jpc,meta.get("Date"),meta.get("kVA"),meta.get("Shift"),meta.get("Phase"),meta.get("Variant"),meta.get("Pole"),obs,dept,station,defect,"Close","",None,""])
    return pd.DataFrame(rows,columns=OBS_HEADERS), pd.DataFrame(review)
